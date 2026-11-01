#!/usr/bin/env python3
"""Helper for the job-search skill: reads and writes tracker.xlsx.

Usage:
  tracker.py init
  tracker.py add '<json object or array>'   # keys = column names; skips duplicates
  tracker.py exists <url> [company] [title]
  tracker.py list [--status "To Review"] [--track Cloud] [--json]
  tracker.py update <row> '<json object>'   # row = spreadsheet row number
"""
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

PATH = Path(__file__).resolve().parent / "tracker.xlsx"
SHEET = "Applications"
COLUMNS = ["Date Found", "Company", "Title", "Location/Remote", "Salary", "Seniority",
           "Match (1-5)", "Why It Fits", "Link", "Source", "Status", "Tailored Resume",
           "Date Applied", "Notes", "Track"]
STATUSES = ["To Review", "Approved", "Applied", "Interviewing", "Offer", "Rejected", "Skipped"]
WIDTHS = [12, 22, 34, 20, 18, 12, 11, 50, 45, 14, 14, 40, 13, 40, 10]
try:
    from search_config import TRACKS as _TRACKS

    TRACKS = list(_TRACKS)
except ImportError:  # tracker.py also works on its own
    TRACKS = ["SWE", "Cloud", "Cyber"]


def init():
    if PATH.exists():
        wb = load_workbook(PATH)
        migrate(wb[SHEET])
        return wb
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(COLUMNS)
    for i, cell in enumerate(ws[1]):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[cell.column_letter].width = WIDTHS[i]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLUMNS)).column_letter}1"
    col = ws.cell(1, COLUMNS.index("Status") + 1).column_letter
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    dv.add(f"{col}2:{col}5000")
    ws.add_data_validation(dv)
    wb.save(PATH)
    return wb


def migrate(ws):
    """Append any columns added since the sheet was created (existing rows get the last track)."""
    header = [c.value for c in ws[1]]
    for i, col in enumerate(COLUMNS):
        if col in header:
            continue
        c = len(header) + 1
        cell = ws.cell(1, c, col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        ws.column_dimensions[cell.column_letter].width = WIDTHS[i]
        if col == "Track":
            for r in range(2, ws.max_row + 1):
                ws.cell(r, c, TRACKS[-1])
        header.append(col)
        ws.parent.save(PATH)


def rows(ws):
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c + 1).value for c in range(len(COLUMNS))]
        if any(v not in (None, "") for v in vals):
            yield r, dict(zip(COLUMNS, vals))


def norm(s):
    return " ".join(str(s or "").lower().split())


def find(ws, url=None, company=None, title=None):
    for r, row in rows(ws):
        if url and norm(row["Link"]).rstrip("/") == norm(url).rstrip("/"):
            return r
        if company and title and norm(row["Company"]) == norm(company) and norm(row["Title"]) == norm(title):
            return r
    return None


def add(payload):
    wb = init()
    ws = wb[SHEET]
    items = payload if isinstance(payload, list) else [payload]
    added, skipped = [], []
    for item in items:
        bad = set(item) - set(COLUMNS)
        if bad:
            sys.exit(f"Unknown columns: {sorted(bad)}")
        dup = find(ws, item.get("Link"), item.get("Company"), item.get("Title"))
        if dup:
            skipped.append({"row": dup, "Company": item.get("Company"), "Title": item.get("Title")})
            continue
        item.setdefault("Date Found", date.today().isoformat())
        item.setdefault("Status", "To Review")
        item.setdefault("Track", TRACKS[-1])
        ws.append([item.get(c) for c in COLUMNS])
        r = ws.max_row
        link = item.get("Link")
        if link:
            cell = ws.cell(r, COLUMNS.index("Link") + 1)
            cell.hyperlink = link
            cell.font = Font(color="0563C1", underline="single")
        for c in range(1, len(COLUMNS) + 1):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
        added.append({"row": r, "Company": item.get("Company"), "Title": item.get("Title")})
    wb.save(PATH)
    print(json.dumps({"added": added, "skipped_duplicates": skipped}, indent=2))


def update(row, changes):
    wb = init()
    ws = wb[SHEET]
    for k, v in changes.items():
        if k not in COLUMNS:
            sys.exit(f"Unknown column: {k}")
        if k == "Track" and v not in TRACKS:
            sys.exit(f"Track must be one of {TRACKS}")
        if k == "Status" and v not in STATUSES:
            sys.exit(f"Status must be one of {STATUSES}")
        ws.cell(row, COLUMNS.index(k) + 1).value = v
    if changes.get("Status") == "Applied" and not ws.cell(row, COLUMNS.index("Date Applied") + 1).value:
        ws.cell(row, COLUMNS.index("Date Applied") + 1).value = date.today().isoformat()
    wb.save(PATH)
    print(json.dumps({"row": row, **dict(rows_by_num(ws, row))}, indent=2, default=str))


def rows_by_num(ws, row):
    return {c: ws.cell(row, i + 1).value for i, c in enumerate(COLUMNS)}.items()


def main():
    args = sys.argv[1:]
    if not args:
        sys.exit(__doc__)
    cmd = args[0]
    if cmd == "init":
        init()
        print(f"Tracker ready: {PATH}")
    elif cmd == "add":
        add(json.loads(args[1]))
    elif cmd == "exists":
        ws = init()[SHEET]
        r = find(ws, args[1] if len(args) > 1 else None, *(args[2:4] if len(args) >= 4 else (None, None)))
        print(r or "")
        sys.exit(0 if r else 1)
    elif cmd == "list":
        ws = init()[SHEET]
        status = args[args.index("--status") + 1] if "--status" in args else None
        track = args[args.index("--track") + 1] if "--track" in args else None
        out = [{"row": r, **row} for r, row in rows(ws)
               if (not status or row["Status"] == status) and (not track or row["Track"] == track)]
        if "--json" in args:
            print(json.dumps(out, indent=2, default=str))
        else:
            for o in out:
                print(f'{o["row"]:>4}  {o["Status"] or "":<12} {o["Match (1-5)"] or "-"}  '
                      f'[{o["Track"] or "-"}] {o["Company"]} — {o["Title"]}  ({o["Location/Remote"] or "?"})  {o["Link"] or ""}')
            print(f"{len(out)} row(s)")
    elif cmd == "update":
        update(int(args[1]), json.loads(args[2]))
    else:
        sys.exit(__doc__)


if __name__ == "__main__":
    main()
